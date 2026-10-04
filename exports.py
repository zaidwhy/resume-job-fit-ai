"""Download builders for the analysis: plain text and Word documents.

Pure functions over the analyzer dataclasses (no Streamlit), so they are unit-tested.
"""

import io

from analyzer import (
    Analysis, CoverLetter, EmailTemplates, InterviewPrep, LinkedInProfile, SkillsRoadmap,
)


def analysis_as_text(
    result: Analysis,
    cover: CoverLetter | None = None,
    prep: InterviewPrep | None = None,
    roadmap: SkillsRoadmap | None = None,
    linkedin: LinkedInProfile | None = None,
    emails: EmailTemplates | None = None,
) -> str:
    lines = [
        "RESUME JOB-FIT ANALYSIS",
        "=" * 40,
        f"Score: {result.score}/100",
        f"Verdict: {result.verdict}",
    ]
    if result.salary_range:
        lines += [f"Estimated Salary Range: {result.salary_range}"]
    lines += [
        "",
        "MATCHED KEYWORDS",
        ", ".join(result.matched_keywords) or "None",
        "",
        "MISSING KEYWORDS",
        ", ".join(result.missing_keywords) or "None",
        "",
        "BULLET REWRITES",
    ]
    for rw in result.bullet_rewrites:
        lines += [f"  Before: {rw.original}", f"  After:  {rw.improved}", ""]
    lines += ["", "HOW TO IMPROVE", result.summary, "", "ATS TIPS"]
    lines += [f"  - {tip}" for tip in result.ats_tips]
    if cover:
        lines += ["", "=" * 40, "COVER LETTER", "=" * 40, "",
                  cover.opening, "", cover.body, "", cover.closing]
    if prep:
        lines += ["", "=" * 40, "INTERVIEW PREP", "=" * 40, ""]
        lines += [f"Key advice: {prep.opening_tip}", ""]
        for i, q in enumerate(prep.questions, 1):
            lines += [f"Q{i}: {q.question}", f"    Why asked: {q.why_asked}",
                      f"    Tip: {q.tip}", ""]
    if roadmap:
        lines += ["", "=" * 40, "SKILLS ROADMAP", "=" * 40, "",
                  f"Timeline: {roadmap.timeline}", "",
                  "Quick wins this week:"]
        lines += [f"  - {w}" for w in roadmap.quick_wins]
        lines += ["", "Skill gaps (priority order):"]
        for gap in roadmap.gaps:
            lines += [f"\n  [{gap.importance}] {gap.skill}", f"  {gap.how_to_learn}"]
            for r in gap.resources:
                lines += [f"    - {r.name} ({r.provider}, {r.type})"]
    if linkedin:
        lines += ["", "=" * 40, "LINKEDIN PROFILE", "=" * 40, "",
                  f"Headline: {linkedin.headline}", "",
                  "ABOUT SECTION", linkedin.about, "",
                  "SKILLS TO ADD", ", ".join(linkedin.skills_to_add), "",
                  "PROFILE TIPS"]
        lines += [f"  - {tip}" for tip in linkedin.profile_tips]
    if emails:
        lines += ["", "=" * 40, "EMAIL TEMPLATES", "=" * 40, "",
                  "--- APPLICATION FOLLOW-UP ---", emails.follow_up, "",
                  "--- POST-INTERVIEW THANK YOU ---", emails.thank_you, "",
                  "--- REJECTION RESPONSE ---", emails.rejection_response]
    return "\n".join(lines)


def export_docx(
    result: Analysis,
    cover: CoverLetter | None = None,
    prep: InterviewPrep | None = None,
    roadmap: SkillsRoadmap | None = None,
    linkedin: LinkedInProfile | None = None,
    emails: EmailTemplates | None = None,
) -> bytes:
    """Build a formatted Word document and return it as bytes."""
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    # Title
    title = doc.add_heading("Resume Job-Fit Analysis", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Score
    score_para = doc.add_paragraph()
    score_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = score_para.add_run(f"Fit Score: {result.score}/100")
    run.bold = True
    run.font.size = Pt(18)
    doc.add_paragraph(result.verdict)
    if result.salary_range:
        p = doc.add_paragraph()
        p.add_run("Estimated Salary Range: ").bold = True
        p.add_run(result.salary_range)

    # Keywords
    doc.add_heading("Matched Keywords", 2)
    doc.add_paragraph(", ".join(result.matched_keywords) or "None")

    doc.add_heading("Missing Keywords", 2)
    doc.add_paragraph(", ".join(result.missing_keywords) or "None")

    # Bullet rewrites
    doc.add_heading("Tailored Bullet Rewrites", 2)
    for rw in result.bullet_rewrites:
        p = doc.add_paragraph()
        p.add_run("Before: ").bold = True
        p.add_run(rw.original)
        p2 = doc.add_paragraph()
        p2.add_run("After:  ").bold = True
        p2.add_run(rw.improved)
        doc.add_paragraph()

    # Summary + ATS
    doc.add_heading("How to Improve This Application", 2)
    doc.add_paragraph(result.summary)
    doc.add_heading("ATS Tips", 3)
    for tip in result.ats_tips:
        doc.add_paragraph(tip, style="List Bullet")

    # Cover letter
    if cover:
        doc.add_page_break()
        doc.add_heading("Cover Letter", 1)
        doc.add_paragraph(cover.opening)
        doc.add_paragraph(cover.body)
        doc.add_paragraph(cover.closing)

    # Interview prep
    if prep:
        doc.add_page_break()
        doc.add_heading("Interview Prep", 1)
        doc.add_paragraph(prep.opening_tip)
        for i, q in enumerate(prep.questions, 1):
            doc.add_heading(f"Q{i}: {q.question}", 3)
            p = doc.add_paragraph()
            p.add_run("Why asked: ").bold = True
            p.add_run(q.why_asked)
            p2 = doc.add_paragraph()
            p2.add_run("Tip: ").bold = True
            p2.add_run(q.tip)

    # Skills roadmap
    if roadmap:
        doc.add_page_break()
        doc.add_heading("Skills Gap Roadmap", 1)
        doc.add_paragraph(f"Estimated timeline: {roadmap.timeline}")
        doc.add_heading("Quick Wins This Week", 2)
        for win in roadmap.quick_wins:
            doc.add_paragraph(win, style="List Bullet")
        doc.add_heading("Skill Gaps - Priority Order", 2)
        for gap in roadmap.gaps:
            doc.add_heading(f"[{gap.importance}] {gap.skill}", 3)
            doc.add_paragraph(gap.how_to_learn)
            for r in gap.resources:
                doc.add_paragraph(f"{r.name} - {r.provider} ({r.type})", style="List Bullet")

    # LinkedIn
    if linkedin:
        doc.add_page_break()
        doc.add_heading("LinkedIn Profile Optimizer", 1)
        doc.add_heading("Headline", 2)
        doc.add_paragraph(linkedin.headline)
        doc.add_heading("About Section", 2)
        doc.add_paragraph(linkedin.about)
        doc.add_heading("Skills to Add", 2)
        doc.add_paragraph(", ".join(linkedin.skills_to_add))
        doc.add_heading("Profile Tips", 2)
        for tip in linkedin.profile_tips:
            doc.add_paragraph(tip, style="List Bullet")

    # Email templates
    if emails:
        doc.add_page_break()
        doc.add_heading("Email Templates", 1)
        for label, body in [
            ("Application Follow-up", emails.follow_up),
            ("Post-Interview Thank You", emails.thank_you),
            ("Rejection Response", emails.rejection_response),
        ]:
            doc.add_heading(label, 2)
            doc.add_paragraph(body)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def export_tailored_resume_docx(result: Analysis) -> bytes:
    """Build a job-tailored resume .docx with AI rewrites substituted for original bullets."""
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    # Contact header placeholder
    name_para = doc.add_paragraph()
    name_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name_run = name_para.add_run("[YOUR NAME]")
    name_run.bold = True
    name_run.font.size = Pt(20)

    contact_para = doc.add_paragraph()
    contact_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact_para.add_run(
        "[your.email@example.com]  ·  [(555) 000-0000]  ·  [linkedin.com/in/yourprofile]  ·  [github.com/yourusername]"
    )

    doc.add_paragraph()

    # Professional summary from AI analysis
    doc.add_heading("Professional Summary", 1)
    doc.add_paragraph(result.summary)

    # Core skills - matched keywords this resume already has
    if result.matched_keywords:
        doc.add_heading("Core Skills", 1)
        doc.add_paragraph(", ".join(result.matched_keywords))

    # Work experience with AI-improved bullets
    if result.bullet_rewrites:
        doc.add_heading("Work Experience", 1)
        exp_header = doc.add_paragraph()
        exp_header.add_run("[Company Name]").bold = True
        exp_header.add_run("  |  [Your Title]  |  [Start Date] – [End Date]")
        for rw in result.bullet_rewrites:
            doc.add_paragraph(rw.improved, style="List Bullet")

    # Skills to add / close gaps
    if result.missing_keywords:
        doc.add_paragraph()
        doc.add_heading("Skills to Develop / Highlight", 1)
        skills_note = doc.add_paragraph()
        skills_note.add_run("Add projects, coursework, or certifications covering: ").italic = True
        skills_note.add_run(", ".join(result.missing_keywords) + ".")

    # ATS tips as an editor's note
    if result.ats_tips:
        doc.add_heading("ATS Optimization Notes", 1)
        ats_note = doc.add_paragraph()
        ats_note.add_run(
            "Before submitting, incorporate the following phrases or sections for better ATS pass-through:"
        ).italic = True
        for tip in result.ats_tips:
            doc.add_paragraph(tip, style="List Bullet")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
